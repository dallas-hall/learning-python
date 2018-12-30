#!/usr/bin/python3

# https://openpyxl.readthedocs.io/en/stable/tutorial.html#loading-from-a-file
from openpyxl import load_workbook
workbook = load_workbook('../Examples/example.xlsx')
print(workbook.sheetnames)
